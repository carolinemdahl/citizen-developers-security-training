"""
Excel-eksporter for Pilar 2 rapporten.

Genererer en Excel-arbeidsbok med 4 faner:
  - Info       : Rapportinformasjon, innholdsbeskrivelse, fargekoder, metodikk
  - Fane 1     : Utbytte (konklusjon) med FRITATT/DELVIS_FRITATT/... fargekoding
  - Fane 2     : LIFO-kjede (audit trail) med lot-status fargekoding
  - Fane 3     : Gevinst/tap salg med gevinst/tap fargekoding
"""

from __future__ import annotations

import io
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

# ---------------------------------------------------------------------------
# Fargepalette
# ---------------------------------------------------------------------------

# Fane 1 – Fritaksstatus
_FILL = {
    "FRITATT":         PatternFill("solid", fgColor="C6EFCE"),   # grønn
    "DELVIS_FRITATT":  PatternFill("solid", fgColor="FFEB9C"),   # gul
    "IKKE_FRITATT":    PatternFill("solid", fgColor="FFC7CE"),   # rød
    "IKKE_EIER":       PatternFill("solid", fgColor="D9D9D9"),   # grå

    # Fane 2 – Lot-status
    "ÅpningsBeholdning":  PatternFill("solid", fgColor="BDD7EE"),  # blå
    "OVERLEVER":           PatternFill("solid", fgColor="C6EFCE"),  # grønn
    "DELVIS_OVERLEVER":    PatternFill("solid", fgColor="FFEB9C"),  # gul
    "KONSUMERT":           PatternFill("solid", fgColor="FFC7CE"),  # rød
    "SALG":                PatternFill("solid", fgColor="E2EFDA"),  # lys grønn (nøytral)
    "UTBYTTE (Dividend)":  PatternFill("solid", fgColor="9DC3E6"),  # blå oppsummering

    # Fane 3 – Gevinst/tap
    "Gevinst":  PatternFill("solid", fgColor="C6EFCE"),
    "Tap":      PatternFill("solid", fgColor="FFC7CE"),

    # Header-farger
    "header_dark":  PatternFill("solid", fgColor="1F4E79"),   # mørk blå
    "header_mid":   PatternFill("solid", fgColor="2E75B6"),   # medium blå
    "header_light": PatternFill("solid", fgColor="D6E4F7"),   # lys blå
    "info_section": PatternFill("solid", fgColor="F2F2F2"),   # lys grå
}

_FONT_WHITE_BOLD = Font(color="FFFFFF", bold=True, size=10)
_FONT_BOLD = Font(bold=True, size=10)
_FONT_NORMAL = Font(size=10)
_FONT_ITALIC = Font(italic=True, size=9, color="595959")

_THIN_BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
_THICK_BOTTOM = Border(bottom=Side(style="medium", color="1F4E79"))


# ---------------------------------------------------------------------------
# Hjelpere
# ---------------------------------------------------------------------------

def _set_col_widths(ws, widths: dict[int, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def _write_header_row(ws, row_num: int, headers: list[str], fill_key: str = "header_mid") -> None:
    fill = _FILL[fill_key]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col, value=header)
        cell.fill = fill
        cell.font = _FONT_WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN_BORDER


def _auto_width(ws, min_w: int = 8, max_w: int = 40) -> None:
    """Autojuster kolonnebredder basert på innhold."""
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(min_w, min(max_len + 3, max_w))


# ---------------------------------------------------------------------------
# Info-fane
# ---------------------------------------------------------------------------

def _write_info_sheet(
    ws,
    portfolio_group: str,
    year: int,
    uttaksdato: date,
) -> None:
    ws.title = "Info"
    ws.sheet_view.showGridLines = False

    def _section(row, title):
        cell = ws.cell(row=row, column=1, value=title)
        cell.font = Font(bold=True, size=11, color="1F4E79")
        cell.border = _THICK_BOTTOM
        ws.merge_cells(f"A{row}:D{row}")

    def _row(row, label, value, note=""):
        ws.cell(row=row, column=2, value=label).font = _FONT_BOLD
        cell_v = ws.cell(row=row, column=3, value=value)
        cell_v.font = _FONT_NORMAL
        if note:
            ws.cell(row=row, column=4, value=note).font = _FONT_ITALIC

    def _color_row(row, label, value, desc, fill_key):
        ws.cell(row=row, column=3, value=label).fill = _FILL[fill_key]
        ws.cell(row=row, column=3).font = _FONT_BOLD
        ws.cell(row=row, column=2, value="").fill = _FILL["info_section"]
        ws.cell(row=row, column=4, value=desc).font = _FONT_ITALIC
        ws.cell(row=row, column=4).alignment = Alignment(wrap_text=True)

    # --- Seksjon 1: Rapportinformasjon ---
    r = 2
    _section(r, "1. Rapportinformasjon")
    _row(r + 1, "Rapporteringsperiode", f"Regnskapsåret {year} (01.01.{year} – 31.12.{year})")
    _row(r + 2, "Porteføljegruppe", portfolio_group)
    _row(r + 3, "Hjemmel",
         "Suppleringsskatteloven § 3-2 første ledd bokstav b (fritatt utbytte) og bokstav c (fritatt egenkapitalgevinst/-tap)")
    _row(r + 4, "Uttaksdato", uttaksdato.strftime("%Y-%m-%d"), "(Snowflake-data per denne dato)")

    # --- Seksjon 2: Innhold per fane ---
    r = 9
    _section(r, "2. Innhold per fane")
    _row(r + 1, "Fane 1 – Utbytte (konklusjon)",
         "Én rad per (portefølje, ISIN, utbetalingsdato). Viser LIFO-overlevende beholdning, andel med eiertid > 12 måneder, fritaksstatus og estimert fritatt beløp NOK.")
    _row(r + 2, "Fane 2 – LIFO-kjede (audit trail)",
         "Detaljert LIFO-kjede per utbyttehendelse. Viser alle kjøp, salg og åpningsbeholdning i 13-månedersvinduet, med lotstatus og akkumulert beholdning.")
    _row(r + 3, "Fane 3 – Gevinst-tap salg",
         "Alle aksjesalg i porteføljegruppen i rapporteringsåret. Inneholder P/L fordelt på kurskomponent og valutakomponent.")

    # --- Seksjon 3: Fargekoder ---
    r = 15
    _section(r, "3. Fargekoder")
    ws.cell(row=r + 1, column=2, value="Fane 1 – Fritaksstatus").font = _FONT_BOLD
    _color_row(r + 2, "FRITATT", "FRITATT", "Hele beholdningen har eiertid ≥ 12 måneder.", "FRITATT")
    _color_row(r + 3, "DELVIS_FRITATT", "DELVIS_FRITATT", "Deler av beholdningen har eiertid ≥ 12 måneder.", "DELVIS_FRITATT")
    _color_row(r + 4, "IKKE_FRITATT", "IKKE_FRITATT", "Ingen del av beholdningen har eiertid ≥ 12 måneder.", "IKKE_FRITATT")
    _color_row(r + 5, "IKKE_EIER", "IKKE_EIER", "Porteføljen hadde ingen LIFO-overlevende beholdning på utbetalingstidspunktet.", "IKKE_EIER")

    ws.cell(row=r + 7, column=2, value="Fane 2 – LIFO-lotstatus").font = _FONT_BOLD
    _color_row(r + 8, "ÅpningsBeholdning", "ÅpningsBeholdning", "Beholdning ved 13-månedersvinduets start. Alltid fritatt.", "ÅpningsBeholdning")
    _color_row(r + 9, "OVERLEVER", "OVERLEVER", "Lot overlever til utbetalingsdato.", "OVERLEVER")
    _color_row(r + 10, "DELVIS_OVERLEVER", "DELVIS_OVERLEVER", "Lot er delvis konsumert av salg.", "DELVIS_OVERLEVER")
    _color_row(r + 11, "KONSUMERT", "KONSUMERT", "Lot er fullt konsumert av salg (LIFO).", "KONSUMERT")
    _color_row(r + 12, "SALG", "SALG", "Salgstransaksjon.", "SALG")
    _color_row(r + 13, "UTBYTTE (Dividend)", "UTBYTTE (Dividend)", "Oppsummeringsrad for utbyttehendelsen.", "UTBYTTE (Dividend)")

    # --- Seksjon 4: Metodikk ---
    r = 32
    _section(r, "4. Metodikk og beregningslogikk")
    metodikk = [
        ("LIFO-prinsipp", "Salgstransaksjoner konsumerer de senest anskaffede lotene først, jf. suppleringsskatteforskriften § 3-2-1 bokstav b nr. 1."),
        ("13-månedersvindu", "For hvert utbytteevent beregnes et rullerende vindu på 395 dager bakover fra utbetalingsdatoen."),
        ("Eiertidsvurdering", "365 dager (≈ 12 måneder) benyttes som terskel for fritak. Åpningsbeholdning er alltid fritatt."),
        ("Fritatt andel", "Fritatt andel = LIFO-overlevende antall med eiertid ≥ 365 dager / total LIFO-overlevende beholdning."),
        ("Fritatt beløp NOK", "Estimert fritatt beløp = Utbytte NOK × Fritatt andel."),
        ("Eierandel", "Estimert eierandel = LIFO-overlevende antall / Totalt utstedt aksjer (ISSUED_VOLUME fra Snowflake)."),
    ]
    for i, (label, text) in enumerate(metodikk):
        _row(r + 1 + i, label, text)

    # --- Seksjon 5: Forbehold ---
    r = 40
    _section(r, "5. Forbehold og begrensninger")
    forbehold = [
        ("Eierandel", "Estimert eierandel er kun for denne porteføljegruppen. Konsernets samlede eierandel kan være høyere."),
        ("Fane 3 – kun realiserte salg", "Fane 3 dekker kun gevinst/tap ved avhendelse. Endringer i virkelig verdi er ikke inkludert."),
        ("Datakilde", "Alle data hentes fra Snowflake DWH_SAM basert på SimCorp Dimension-data."),
    ]
    for i, (label, text) in enumerate(forbehold):
        _row(r + 1 + i, label, text)

    # --- Ansvarsfraskrivelse ---
    ws.cell(row=46, column=1,
            value="Rapporten er maskinelt generert og må gjennomgås og kvalitetssikres av ansvarlig skatterådgiver før bruk i formell rapportering.").font = Font(bold=True, italic=True, size=9, color="9C0006")
    ws.merge_cells("A46:D46")

    # Kolonnebredder og radhøyder
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 55
    ws.column_dimensions["D"].width = 55
    ws.row_dimensions[1].height = 8

    for row_cells in ws.iter_rows(min_row=2, max_row=46, min_col=2, max_col=4):
        for cell in row_cells:
            if not cell.fill.fgColor.rgb or cell.fill.fgColor.rgb == "00000000":
                cell.fill = _FILL["info_section"]
            cell.alignment = Alignment(wrap_text=True, vertical="top")


# ---------------------------------------------------------------------------
# Fane 1: Utbytte konklusjon
# ---------------------------------------------------------------------------

def _write_fane1(ws, df: pd.DataFrame) -> None:
    ws.title = "1 - Utbytte (konklusjon)"
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    if df.empty:
        ws.cell(row=1, column=1, value="Ingen data")
        return

    headers = list(df.columns)
    _write_header_row(ws, 1, headers, "header_dark")

    # Finn kolonneindeks for Fritaksstatus
    status_col = headers.index("Fritaksstatus eiertid") + 1 if "Fritaksstatus eiertid" in headers else None
    nok_col = headers.index("Utbytte NOK") + 1 if "Utbytte NOK" in headers else None
    fritatt_col = headers.index("Eiertid > 12 mnd beløp NOK") + 1 if "Eiertid > 12 mnd beløp NOK" in headers else None

    for row_idx, (_, row) in enumerate(df.iterrows(), 2):
        status = str(row.get("Fritaksstatus eiertid", ""))
        fill = _FILL.get(status)

        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = _FONT_NORMAL
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(horizontal="left", vertical="center")

            if col_idx == status_col and fill:
                cell.fill = fill
                cell.font = _FONT_BOLD

            # Høyrejuster tall
            if col_idx in (nok_col, fritatt_col) and isinstance(val, (int, float)):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
            elif isinstance(val, (int, float)) and col_idx not in (status_col,):
                cell.alignment = Alignment(horizontal="right")

    _auto_width(ws)
    ws.row_dimensions[1].height = 30


# ---------------------------------------------------------------------------
# Fane 2: LIFO-kjede (audit trail)
# ---------------------------------------------------------------------------

def _write_fane2(ws, df: pd.DataFrame) -> None:
    ws.title = "2 - LIFO-kjede (audit trail)"
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    if df.empty:
        ws.cell(row=1, column=1, value="Ingen data")
        return

    headers = list(df.columns)
    _write_header_row(ws, 1, headers, "header_dark")

    lot_col = headers.index("Lot-status") + 1 if "Lot-status" in headers else None

    for row_idx, (_, row) in enumerate(df.iterrows(), 2):
        lot_status = str(row.get("Lot-status", ""))
        fill = _FILL.get(lot_status)

        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = _FONT_NORMAL
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(horizontal="left", vertical="center")

            if fill:
                cell.fill = fill
            if col_idx == lot_col:
                cell.font = _FONT_BOLD

            if isinstance(val, (int, float)) and col_idx != lot_col:
                cell.alignment = Alignment(horizontal="right")

    _auto_width(ws)
    ws.row_dimensions[1].height = 30


# ---------------------------------------------------------------------------
# Fane 3: Gevinst/tap salg
# ---------------------------------------------------------------------------

def _write_fane3(ws, df: pd.DataFrame) -> None:
    ws.title = "3 - Gevinst-tap salg"
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    if df.empty:
        ws.cell(row=1, column=1, value="Ingen data")
        return

    headers = list(df.columns)
    _write_header_row(ws, 1, headers, "header_dark")

    pl_col = headers.index("pl_total_nok") + 1 if "pl_total_nok" in headers else None

    for row_idx, (_, row) in enumerate(df.iterrows(), 2):
        pl_val = row.get("pl_total_nok")
        fill = None
        if pl_val is not None:
            fill = _FILL["Gevinst"] if float(pl_val) >= 0 else _FILL["Tap"]

        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = _FONT_NORMAL
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(horizontal="left", vertical="center")

            if col_idx == pl_col and fill:
                cell.fill = fill
                cell.font = _FONT_BOLD

            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right")
                if col_idx in (
                    headers.index(c) + 1
                    for c in ("salgsproveny_nok", "kostverdi_nok",
                               "pl_kurs_nok", "pl_valuta_nok", "pl_total_nok")
                    if c in headers
                ):
                    cell.number_format = "#,##0.00"

    _auto_width(ws)
    ws.row_dimensions[1].height = 30


# ---------------------------------------------------------------------------
# Hoved-eksportfunksjon
# ---------------------------------------------------------------------------

def generate_excel(
    portfolio_group: str,
    year: int,
    uttaksdato: date,
    fane1_df: pd.DataFrame,
    fane2_df: pd.DataFrame,
    fane3_df: pd.DataFrame,
) -> bytes:
    """
    Generer Excel-arbeidsboken og returner som bytes (for Streamlit-nedlasting).
    """
    wb = Workbook()

    # Info-fane
    ws_info = wb.active
    _write_info_sheet(ws_info, portfolio_group, year, uttaksdato)

    # Fane 1
    ws1 = wb.create_sheet()
    _write_fane1(ws1, fane1_df)

    # Fane 2
    ws2 = wb.create_sheet()
    _write_fane2(ws2, fane2_df)

    # Fane 3
    ws3 = wb.create_sheet()
    _write_fane3(ws3, fane3_df)

    # Returner som bytes
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
